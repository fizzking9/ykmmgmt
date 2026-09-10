"""Tests for the raw business-source cleansing pipeline (raw_cleaning).

Unit tests build synthetic "non-standard CSV" files reproducing the
documented source defects (A1–A7, B1–B2) and verify the deterministic
repairs plus idempotency. The integration test runs a dirty raw export
through the real import endpoint against a dynamically created table.
"""

import uuid

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

from app.services.raw_cleaning import RawCleanError, parse_with_profile
from app.services.raw_cleaning.common import (
    align_by_name,
    cell_to_str,
    parse_csv_line,
    read_physical_lines,
    strip_trailing_empty,
)
from app.services.raw_cleaning.tables import (
    get_profile,
    match_profile,
)
from app.services.raw_cleaning.tables import refund_export as refund_export_profile
from app.services.raw_cleaning.tables import service_order as service_order_profile
from tests.schema_cleanup import purge_dynamic_table

# ── service_order (服务工单列表) ────────────────────────────────────────────

SO_HEADER_23 = ",".join(service_order_profile.HEADER) + ","
SO_HEADER_22 = ",".join(c for c in service_order_profile.HEADER if c != "服务人员") + ","

SO_DIRTY_ROW = (
    "GD001\t,SN001\t,洗衣机\t,滚筒,13800000000,安装,上门安装,高,待处理,备注文本,"
    "2026-08-06 10:00:00,2026-08-06 10:05:00,2026-08-06 11:00:00,2026-08-06 12:00:00,"
    "60分钟,张三,李四,APP,节点1,王五,已处理,否,{\"a\":\"1,2\",\"b\":\"x\"},"
)
SO_CLEAN_ROW = (
    "GD002,SN002,烘干机,,13900000000,维修,上门维修,低,已完成,另一条备注,"
    "2026-08-07 10:00:00,2026-08-07 10:05:00,2026-08-07 11:00:00,2026-08-07 12:00:00,"
    "30分钟,赵六,钱七,小程序,节点2,孙八,已完成,是,"
)


def _write_csv(path, content: str, newline: str = "\r\n") -> None:
    # newline="" prevents the platform newline translation that would
    # otherwise corrupt the explicit \r\n / bare-\n byte patterns.
    path.write_text(content.replace("\n", newline), encoding="utf-8", newline="")


class TestServiceOrderProfile:
    def test_unescaped_json_and_trailing_tabs(self, tmp_path):
        f = tmp_path / "so.csv"
        _write_csv(f, SO_HEADER_23 + "\n" + SO_DIRTY_ROW + "\n")

        df = service_order_profile.clean_file(f)
        assert list(df.columns) == service_order_profile.HEADER
        assert len(df) == 1
        row = df.iloc[0]
        assert row["工单号"] == "GD001"  # trailing tab stripped
        assert row["SN"] == "SN001"
        assert row["服务人员"] == "李四"
        assert row["提交参数"] == '{"a":"1,2","b":"x"}'  # JSON recovered whole

    def test_column_evolution_22_to_23(self, tmp_path):
        """Old exports lack 服务人员 — aligned to the 23-column superset."""
        f = tmp_path / "so_old.csv"
        old_row = SO_CLEAN_ROW.replace("钱七,", "")  # drop the 服务人员 value
        _write_csv(f, SO_HEADER_22 + "\n" + old_row + "\n")

        df = service_order_profile.clean_file(f)
        assert list(df.columns) == service_order_profile.HEADER
        row = df.iloc[0]
        assert row["服务人员"] == ""  # missing column filled empty
        assert row["工单号"] == "GD002"
        assert row["登记人"] == "赵六"

    def test_idempotent_on_cleansed_output(self, tmp_path):
        """Re-cleansing a cleansed standard CSV yields the identical result."""
        f1 = tmp_path / "so_raw.csv"
        _write_csv(f1, SO_HEADER_23 + "\n" + SO_DIRTY_ROW + "\n" + SO_CLEAN_ROW + "\n")
        df1 = service_order_profile.clean_file(f1)

        f2 = tmp_path / "so_clean.csv"
        df1.to_csv(f2, index=False)
        df2 = service_order_profile.clean_file(f2)
        assert df1.equals(df2)


# ── refund_export (退费单导出列表) ─────────────────────────────────────────

RE_HEADER = ",".join(refund_export_profile.HEADER) + ","


class TestRefundExportProfile:
    def test_remark_with_half_width_comma(self, tmp_path):
        """A5: commas inside 备注 inflate the field count — repaired via
        the 状态+退款方式 enum pair."""
        f = tmp_path / "re.csv"
        row = (
            "TF001,PO001,TP001,123456,设备坏了,套餐A,商户甲,100.00,95.00,"
            "用户说,要求部分退款,已退款,线下,同意,审核人A,"
            "2026-08-06 09:00:00,2026-08-06 10:00:00,操作人A,30.00,"
        )
        _write_csv(f, RE_HEADER + "\n" + row + "\n")

        df = refund_export_profile.clean_file(f)
        assert len(df) == 1
        r = df.iloc[0]
        assert r["备注"] == "用户说,要求部分退款"
        assert r["状态"] == "已退款"
        assert r["退款方式"] == "线下"
        assert r["审核备注"] == "同意"
        assert r["套餐价格"] == "30.00"

    def test_embedded_bare_lf_in_remark(self, tmp_path):
        """A1: bare LF inside 备注 stays in the field value (physical
        lines split on CRLF)."""
        f = tmp_path / "re_nl.csv"
        row = (
            "TF002,PO002,TP002,123457,设备丢了,套餐B,商户乙,50.00,50.00,"
            "第一行\n第二行,已退款,线下,,审核人B,"
            "2026-08-07 09:00:00,2026-08-07 10:00:00,操作人B,20.00,"
        )
        # Written raw: CRLF physical line endings with a bare LF kept
        # inside the remark field.
        f.write_text(RE_HEADER + "\r\n" + row + "\r\n", encoding="utf-8", newline="")

        df = refund_export_profile.clean_file(f)
        assert len(df) == 1
        assert df.iloc[0]["备注"] == "第一行\n第二行"

    def test_lf_only_file(self, tmp_path):
        """A6: files with plain LF line endings parse row-by-row."""
        f = tmp_path / "re_lf.csv"
        row = (
            "TF003,PO003,TP003,123458,正常,套餐C,商户丙,10.00,10.00,ok,"
            "已退款,线上,ok,审核人C,2026-08-08 09:00:00,2026-08-08 10:00:00,操作人C,5.00,"
        )
        _write_csv(f, RE_HEADER + "\n" + row + "\n", newline="\n")

        df = refund_export_profile.clean_file(f)
        assert len(df) == 1
        assert df.iloc[0]["退费单号"] == "TF003"

    def test_missing_first_column_whole_file_shift(self, tmp_path):
        """A7: exports without the 退费单号 column shift every value one
        left — detected file-wide and shifted back."""
        f = tmp_path / "re_shift.csv"
        # 17 values: 平台订单号..套餐价格 (退费单号 missing entirely)
        rows = []
        for i in range(3):
            rows.append(
                f"PO{i:03d},TP{i:03d},{100000 + i},设备坏了{i},套餐A,商户甲,100.00,95.00,"
                f"备注{i},已退款,转入提现账户,审核备注{i},审核人A,"
                f"2026-08-11 09:0{i}:00,2026-08-11 10:0{i}:00,操作人A,30.00,"
            )
        _write_csv(f, RE_HEADER + "\n" + "\n".join(rows) + "\n")

        df = refund_export_profile.clean_file(f)
        assert len(df) == 3
        for i, (_, r) in enumerate(df.iterrows()):
            assert r["退费单号"] == ""  # restored empty (was missing)
            assert r["平台订单号"] == f"PO{i:03d}"
            assert r["三方订单号"] == f"TP{i:03d}"
            assert r["设备SN"] == str(100000 + i)
            assert r["退费原因"] == f"设备坏了{i}"
            assert r["套餐名称"] == "套餐A"
            assert r["状态"] == "已退款"
            assert r["退款方式"] == "转入提现账户"
            assert r["套餐价格"] == "30.00"  # back at its own position

    def test_shift_detection_not_triggered_on_normal_file(self, tmp_path):
        f = tmp_path / "re_normal.csv"
        row = (
            "TF004,PO004,TP004,123459,正常,套餐D,商户丁,10.00,10.00,ok,"
            "已退款,线上,ok,审核人D,2026-08-09 09:00:00,2026-08-09 10:00:00,操作人D,5.00,"
        )
        _write_csv(f, RE_HEADER + "\n" + row + "\n")

        df = refund_export_profile.clean_file(f)
        assert df.iloc[0]["退费单号"] == "TF004"  # NOT shifted

    def test_multiline_quoted_field_standard_csv(self, tmp_path):
        """Standard-CSV quoted field spanning physical lines (embedded LF
        inside quotes) must reassemble into ONE row with the field whole
        — not fragment into extra rows."""
        f = tmp_path / "re_mlq.csv"
        content = (
            RE_HEADER + "\n"
            'TF007,PO007,TP007,123461,原因,套餐F,商户己,10.00,10.00,"备注第一行\n备注第二行,还带逗号",'
            "已退款,线下,ok,审核人G,2026-08-11 09:00:00,2026-08-11 10:00:00,操作人G,5.00,\n"
            "TF008,PO008,TP008,123462,原因2,套餐G,商户庚,20.00,20.00,普通备注,已退款,线上,ok,审核人H,"
            "2026-08-12 09:00:00,2026-08-12 10:00:00,操作人H,8.00,\n"
        )
        f.write_text(content, encoding="utf-8", newline="")

        df = refund_export_profile.clean_file(f)
        assert len(df) == 2  # not 3 — the continuation line must not fragment
        assert df.iloc[0]["备注"] == "备注第一行\n备注第二行,还带逗号"
        assert df.iloc[0]["退费单号"] == "TF007"
        assert df.iloc[1]["退费单号"] == "TF008"

    def test_idempotent_on_cleansed_output_with_embedded_newline(self, tmp_path):
        """Cleansed output (quoted, multi-line) re-cleanses identically."""
        f1 = tmp_path / "re_raw.csv"
        row = (
            "TF009,PO009,TP009,123463,原因,套餐H,商户辛,100.00,95.00,"
            "备注行一\n备注行二,已退款,线下,审核备注,审核人I,"
            "2026-08-10 09:00:00,2026-08-10 10:00:00,操作人I,30.00,"
        )
        # raw A1-style input: CRLF record endings, bare LF inside 备注
        f1.write_text(RE_HEADER + "\r\n" + row + "\r\n", encoding="utf-8", newline="")
        df1 = refund_export_profile.clean_file(f1)
        assert df1.iloc[0]["备注"] == "备注行一\n备注行二"

        f2 = tmp_path / "re_clean.csv"
        df1.to_csv(f2, index=False)  # pandas quotes the embedded newline
        df2 = refund_export_profile.clean_file(f2)
        assert len(df2) == 1
        assert df1.equals(df2)

    def test_idempotent_on_cleansed_output(self, tmp_path):
        f1 = tmp_path / "re_raw.csv"
        row = (
            "TF005,PO005,TP005,123460,原因,套餐E,商户戊,100.00,95.00,"
            "备注,含逗号,已退款,线下,审核备注,审核人E,"
            "2026-08-10 09:00:00,2026-08-10 10:00:00,操作人E,30.00,"
        )
        _write_csv(f1, RE_HEADER + "\n" + row + "\n")
        df1 = refund_export_profile.clean_file(f1)
        assert df1.iloc[0]["备注"] == "备注,含逗号"

        f2 = tmp_path / "re_clean.csv"
        df1.to_csv(f2, index=False)
        df2 = refund_export_profile.clean_file(f2)
        assert df1.equals(df2)


# ── plain CSV profiles ─────────────────────────────────────────────────────


class TestPlainCsvProfiles:
    def test_withdraw_record_trailing_tabs_and_comma(self, tmp_path):
        profile = get_profile("withdraw_record")
        f = tmp_path / "wd.csv"
        header = ",".join(profile.HEADER) + ","
        row = "69212612184185\t,120.50,SN777\t,提现,50.00\t,备注,2026-08-06 10:00:00,管理员,"
        _write_csv(f, header + "\n" + row + "\n")

        df = profile.clean_file(f)
        assert df.iloc[0]["账户ID"] == "69212612184185"
        assert df.iloc[0]["SN"] == "SN777"
        assert df.iloc[0]["操作金额"] == "50.00"

    def test_device_tag_lowercase_sn_header(self, tmp_path):
        profile = get_profile("标签")
        f = tmp_path / "dt.csv"
        header = ",".join(profile.HEADER) + ","
        row = "69212612184186\t,重要客户,商户A,套餐F,admin,2026-08-06 10:00:00,"
        _write_csv(f, header + "\n" + row + "\n")

        df = profile.clean_file(f)
        assert df.iloc[0]["设备sn"] == "69212612184186"
        assert df.iloc[0]["标签名"] == "重要客户"

    def test_device_package(self, tmp_path):
        profile = get_profile("device_package")
        f = tmp_path / "dp.csv"
        header = ",".join(profile.HEADER) + ","
        row = ("69212612184187\t,套餐G,月套餐,通用,商户B,30.00,2026-08-01,2026-08-31,"
               "2026-08-06 10:00:00,2026-08-01 09:00:00,30,admin,PO123,10240,2048,8192,使用中,")
        _write_csv(f, header + "\n" + row + "\n")

        df = profile.clean_file(f)
        assert df.iloc[0]["设备SN"] == "69212612184187"
        assert df.iloc[0]["套餐有效天数"] == "30"


# ── xlsx profiles ──────────────────────────────────────────────────────────


class TestXlsxProfiles:
    def test_ai_acc_cs_record_value_normalization(self, tmp_path):
        import datetime

        profile = get_profile("ai_acc_cs_record")
        f = tmp_path / "ai.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(profile.HEADER)
        ws.append(["网络", "无法联网", "SN888", "重启", datetime.datetime(2026, 8, 6, 10, 0, 0), "客服A",
                   "请重启设备", "说明", "已解决", "是", None, None, 100.0])
        wb.save(f)

        df = profile.clean_file(f)
        row = df.iloc[0]
        assert row["操作时间"] == "2026-08-06 10:00:00"  # datetime → string
        assert row["最后咨询时间"] == "100"  # integral float → int form
        assert row["咨询时间"] == ""  # None → ""

    def test_refund_info_csv_upload_still_works(self, tmp_path):
        """xlsx-native tables may arrive as CSV — the profile handles both."""
        profile = get_profile("退款单信息")
        f = tmp_path / "ri.csv"
        header = ",".join(profile.HEADER) + ","
        row = "TF006,PO006,13700000000,套餐H,,10.00,10.00,不想要了,已退款,线上,审核人F,2026-08-06 09:00:00,,"
        _write_csv(f, header + "\n" + row + "\n")

        df = profile.clean_file(f)
        assert list(df.columns) == profile.HEADER
        assert df.iloc[0]["手机号码"] == "13700000000"
        assert df.iloc[0]["实际支付金额"] == ""  # never collected — kept empty


# ── common helpers ─────────────────────────────────────────────────────────


class TestCommonHelpers:
    def test_parse_csv_line_quoted_and_escaped(self):
        fields = parse_csv_line('a,"b,c","say ""hi""",d')
        assert fields == ["a", "b,c", 'say "hi"', "d"]

    def test_strip_trailing_empty(self):
        assert strip_trailing_empty(["a", "b", "", ""]) == ["a", "b"]
        assert strip_trailing_empty([]) == []

    def test_read_physical_lines_crlf_and_bom(self, tmp_path):
        f = tmp_path / "bom.csv"
        f.write_bytes(b"\xef\xbb\xbfa,b\r\nc,d\r\n")
        # trailing \r\n leaves one empty tail element (skipped by callers)
        assert read_physical_lines(f)[:2] == ["a,b", "c,d"]
        assert read_physical_lines(f)[-1] == ""

    def test_read_physical_lines_lf_only(self, tmp_path):
        f = tmp_path / "lf.csv"
        f.write_bytes(b"a,b\nc,d\n")
        assert read_physical_lines(f) == ["a,b", "c,d", ""]

    def test_cell_to_str(self):
        import datetime

        assert cell_to_str(None) == ""
        assert cell_to_str(" x ") == "x"
        assert cell_to_str(datetime.datetime(2026, 8, 6, 9, 30, 0)) == "2026-08-06 09:30:00"
        assert cell_to_str(datetime.date(2026, 8, 6)) == "2026-08-06"
        assert cell_to_str(100.0) == "100"
        assert cell_to_str(1.5) == "1.5"
        assert cell_to_str(True) == "TRUE"

    def test_align_by_name_fills_missing(self):
        header = ["a", "b"]
        canonical = ["a", "b", "c"]
        assert align_by_name(header, ["1", "2"], canonical) == ["1", "2", ""]
        assert align_by_name(["b", "a"], ["2", "1"], canonical) == ["1", "2", ""]


# ── profile matching & dispatch ────────────────────────────────────────────


class TestProfileMatching:
    def test_exact_header_set_match(self):
        profile = match_profile(service_order_profile.HEADER)
        assert profile is service_order_profile
        profile = match_profile(refund_export_profile.HEADER)
        assert profile is refund_export_profile

    def test_signature_overlap_match(self):
        # Subset of a device_package export (partial column list)
        assert match_profile(["设备SN", "套餐名称", "套餐类型", "可用流量"]).TABLE == "设备套餐列表"

    def test_generic_headers_no_match(self):
        assert match_profile(["订单号", "金额", "状态", "记录时间"]) is None
        assert match_profile(["col1", "col2"]) is None

    def test_get_profile_aliases(self):
        assert get_profile("工单").TABLE == "服务工单列表"
        assert get_profile("service_order").TABLE == "服务工单列表"
        assert get_profile("退费").TABLE == "退费单导出列表"
        assert get_profile("nope") is None


class TestParseWithProfile:
    def test_passthrough_for_generic_csv(self, tmp_path):
        f = tmp_path / "generic.csv"
        f.write_text("订单号,金额\nA1,10.5\n", encoding="utf-8")
        df, headers, profile = parse_with_profile(f)
        assert profile is None
        assert headers == ["订单号", "金额"]
        assert len(df) == 1

    def test_dirty_service_order_dispatch(self, tmp_path):
        f = tmp_path / "so.csv"
        _write_csv(f, SO_HEADER_23 + "\n" + SO_DIRTY_ROW + "\n")
        df, headers, profile = parse_with_profile(f)
        assert profile == "服务工单列表"
        assert headers == service_order_profile.HEADER
        assert df.iloc[0]["提交参数"] == '{"a":"1,2","b":"x"}'

    def test_corrupt_xlsx_raises(self, tmp_path):
        """B2: an export that is actually an API-error JSON body."""
        f = tmp_path / "corrupt.xlsx"
        f.write_text('{"code": 500, "msg": "internal error"}', encoding="utf-8")
        with pytest.raises(RawCleanError):
            parse_with_profile(f)


# ── integration: dirty raw export through the real import endpoint ────────


@pytest.mark.asyncio
async def test_raw_export_import_end_to_end(tmp_path):
    """A dirty 服务工单列表 export flows through the raw cleanse and lands
    correctly in a dynamically created table."""
    name = f"raw_so_{uuid.uuid4().hex[:6]}"
    from main import app

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        try:
            # Table whose column labels equal the canonical Chinese headers
            str_cols = [
                ("order_no", "工单号"), ("sn", "SN"), ("device_type", "设备类型"),
                ("device_type_remark", "设备类型备注"), ("phone", "手机号"),
                ("service_category", "服务类别"), ("service_item", "服务项"),
                ("priority", "优先级"), ("status", "状态"), ("handle_duration", "处理时长"),
                ("registrar", "登记人"), ("service_staff", "服务人员"),
                ("entry_channel", "录入渠道"), ("handle_node", "处理节点"),
                ("handler", "处理人"), ("is_appealed", "是否申诉"),
            ]
            dt_cols = [
                ("register_time", "登记时间"), ("activate_time", "激活时间"),
                ("dispatch_time", "分发时间"), ("finish_time", "完成时间"),
            ]
            text_cols = [("remark", "备注"), ("handle_opinion", "处理意见"), ("submit_params", "提交参数")]
            columns = [
                {"name": c, "type": "String", "length": 100, "nullable": True, "label": lbl}
                for c, lbl in str_cols
            ] + [
                {"name": c, "type": "DateTime", "nullable": True, "label": lbl} for c, lbl in dt_cols
            ] + [
                {"name": c, "type": "Text", "nullable": True, "label": lbl} for c, lbl in text_cols
            ]
            resp = await client.post(
                "/api/schema/tables", json={"name": name, "display_name": "服务工单导入测试", "columns": columns}
            )
            assert resp.status_code == 201, resp.text

            dirty = tmp_path / "服务工单列表_20260806.csv"
            _write_csv(dirty, SO_HEADER_23 + "\n" + SO_DIRTY_ROW + "\n" + SO_CLEAN_ROW + "\n")
            resp = await client.post(
                "/api/imports",
                files={"file": (dirty.name, dirty.read_bytes(), "text/csv")},
                data={"target_table": name},
            )
            assert resp.status_code == 200, resp.text
            body = resp.json()
            assert body["rows_inserted"] == 2
            assert body["rows_rejected"] == 0
            step_names = [s["step"] for s in body["cleaning_report"]["steps"]]
            assert "raw_source_cleanse" in step_names

            rows = (await client.get(f"/api/schema/tables/{name}")).json()["sample_rows"]
            by_order = {r["order_no"]: r for r in rows}
            assert by_order["GD001"]["sn"] == "SN001"  # tab stripped
            assert by_order["GD001"]["submit_params"] == '{"a":"1,2","b":"x"}'
            assert by_order["GD001"]["service_staff"] == "李四"
            assert by_order["GD002"]["service_staff"] == "钱七"
        finally:
            await purge_dynamic_table(name)

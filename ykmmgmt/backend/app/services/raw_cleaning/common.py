"""Shared low-level utilities for raw business-source cleansing.

The source system's exports are "non-standard CSV": the bytes are
self-consistent, but the exporter does not follow standard CSV escaping
rules, so a standard parser misreads them (fields shift, rows split).
These helpers implement the correct re-parsing rules (numbering follows
the data-quality report of the source exports):

  A1/A6  physical lines end with CRLF while newlines *inside* field
         values are bare LF → split on CRLF so bare LF stays inside the
         field value; files with plain LF line endings throughout are
         split on LF instead (otherwise they parse as a single line);
  A3     some field values carry a trailing tab → strip();
  A4     every line ends with one extra comma (one extra empty field) →
         drop trailing empty fields;
  A2/A5  fields are comma-separated, occasionally wrapped in double
         quotes → hand-written parser that also handles the quoted form
         ("" escapes), so cleansing is idempotent on already-cleaned
         standard CSV files.
"""

from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


def _decode(data: bytes) -> str:
    """Decode raw bytes: UTF-8 first, GBK fallback, then lossy UTF-8."""
    for enc in ("utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def read_physical_lines(path: Path | str) -> list[str]:
    """Read a CSV's raw bytes, drop the BOM, split into logical lines.

    - Most files end lines with CRLF (``\\r\\n``) while field values
      contain bare LF (``\\n``) → split on ``\\r\\n`` so the bare LF stays
      inside the field value.
    - A few files use plain LF line endings throughout → split on ``\\n``.
    - Standard-CSV quoted fields may span multiple physical lines
      (embedded newline inside a quoted field). A quote-state scan
      rejoins those continuation lines with ``\\n`` so the field stays
      whole; unquoted dirty rows (never inside quotes at end-of-line)
      are unaffected, keeping the A1 bare-LF behavior intact.
    """
    data = open(path, "rb").read()
    if data[:3] == b"\xef\xbb\xbf":
        data = data[3:]
    text = _decode(data)
    sep = "\r\n" if "\r\n" in text else "\n"
    physical = text.split(sep)

    lines: list[str] = []
    buf: str | None = None
    for ln in physical:
        buf = ln if buf is None else buf + "\n" + ln
        if _in_open_quote(buf):
            continue  # line ends inside a quoted field — keep joining
        lines.append(buf)
        buf = None
    if buf is not None:
        lines.append(buf)  # unterminated quote — emit what we have
    return lines


def _in_open_quote(s: str) -> bool:
    """Whether the line ends inside an open double-quoted field
    (odd number of quotes not consumed by "" escapes)."""
    in_q = False
    i = 0
    n = len(s)
    while i < n:
        c = s[i]
        if c == '"':
            if in_q and i + 1 < n and s[i + 1] == '"':
                i += 1  # escaped "" inside quotes
            else:
                in_q = not in_q
        i += 1
    return in_q


def parse_csv_line(s: str) -> list[str]:
    """Hand-written CSV parse: comma delimiter + double-quote wrapping
    ("" escape) + bare LF/CR kept as ordinary characters."""
    fields = []
    cur = []
    in_q = False
    i = 0
    n = len(s)
    while i < n:
        c = s[i]
        if in_q:
            if c == '"':
                if i + 1 < n and s[i + 1] == '"':
                    cur.append('"')
                    i += 1
                else:
                    in_q = False
            else:
                cur.append(c)
        else:
            if c == '"':
                in_q = True
            elif c == ",":
                fields.append("".join(cur))
                cur = []
            else:
                cur.append(c)  # bare LF / CR kept inside the field value
        i += 1
    fields.append("".join(cur))
    return fields


def clean_cell(v: str) -> str:
    """Strip leading/trailing whitespace (tabs, full-width spaces, etc.)."""
    return v.strip()


def strip_trailing_empty(fields: list[str]) -> list[str]:
    """Drop the empty fields produced by the trailing extra comma."""
    while fields and fields[-1] == "":
        fields.pop()
    return fields


def cell_to_str(v) -> str:
    """Normalize an xlsx cell value to a string:
    None→"", str→strip, datetime→string, integral float→int form."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return repr(v)
    return str(v)


def iter_xlsx_rows(path: Path | str):
    """Iterate an xlsx's rows read-only (values stringified, trailing
    empties dropped, fully blank rows skipped)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            row = [cell_to_str(v) for v in r]
            while row and row[-1] == "":
                row.pop()
            if any(x != "" for x in row):
                yield row
    finally:
        wb.close()


def read_csv_rows(file: Path | str) -> tuple[list[str], list[list[str]]]:
    """Read a CSV file → (header, rows).

    Rows get the common cleanse (strip + trailing empty drop). Suitable
    for tables without special structure (no embedded JSON / newlines).
    """
    lines = read_physical_lines(file)
    header = [clean_cell(x) for x in parse_csv_line(lines[0])]
    rows = []
    for s in lines[1:]:
        if s.strip() == "":
            continue
        fields = [clean_cell(x) for x in parse_csv_line(s)]
        fields = strip_trailing_empty(fields)
        rows.append(fields)
    return header, rows


def read_table_rows(file: Path | str) -> tuple[list[str], list[list[str]]]:
    """Read a table file by extension → (header, rows).

    - .xlsx/.xlsm → openpyxl read-only (values normalized via cell_to_str)
    - otherwise   → CSV parse (strip + trailing empty drop)

    The same table can thus be read either from the original Excel export
    or from a cleansed standard CSV, and the value normalization is
    idempotent (datetime→string and integral float→int are stable on a
    second pass).
    """
    if str(file).lower().endswith((".xlsx", ".xlsm")):
        it = iter_xlsx_rows(file)
        header = list(next(it))
        return header, list(it)
    return read_csv_rows(file)


def align_by_name(header: list[str], row: list[str], canonical: list[str]) -> list[str]:
    """Align a row to the canonical column order by header name; columns
    missing from the file are filled with "" (column evolution / reordering)."""
    idx: dict[str, int] = {}
    for i, c in enumerate(header):
        c = (c or "").strip()
        idx.setdefault(c, i)
    out = [""] * len(canonical)
    for j, c in enumerate(canonical):
        k = idx.get(c)
        if k is not None and k < len(row):
            out[j] = row[k]
    return out
